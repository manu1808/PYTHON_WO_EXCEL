"""
Lógica para la depuración y formateo del archivo RAW de Work Orders (WO).
"""

import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from config import (
    COLUMNS_TO_KEEP,
    COL_WIDTHS,
    HEADER_ROW_HEIGHT,
    DATA_ROW_HEIGHT,
    NEG_FILL,
    NEG_COLOR
)

class ExcelTransformer:
    @staticmethod
    def transform(raw_path: str, out_path: str, progress_cb=None) -> int:
        """
        Limpia, filtra columnas, formatea y resalta valores negativos en el Excel RAW.
        Retorna la cantidad de celdas negativas resaltadas.
        """
        def progress(msg):
            if progress_cb:
                progress_cb(msg)

        progress("Preparando archivo temporal…")
        shutil.copy2(raw_path, out_path)

        progress("Abriendo workbook…")
        wb = load_workbook(out_path)
        ws = wb.active

        # Deshacer celdas combinadas
        for merge in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merge))

        # Buscar fila de encabezado
        header_row_1based = None
        for row in ws.iter_rows():
            if row[0].value == 'No.':
                header_row_1based = row[0].row
                break

        if header_row_1based is None:
            raise ValueError(
                "No se encontró la fila de encabezado ('No.') en el archivo.\n"
                "Verifica que sea el export RAW correcto."
            )

        # Mapa de columnas inicial
        col_name_to_idx = {}
        for cell in ws[header_row_1based]:
            if cell.value:
                col_name_to_idx[str(cell.value).strip()] = cell.column

        missing = set(COLUMNS_TO_KEEP) - set(col_name_to_idx.keys())
        if missing:
            raise ValueError(
                "Columnas no encontradas en el RAW:\n" +
                "\n".join(f"- {column}" for column in missing)
            )

        # Eliminar columnas innecesarias
        progress("Eliminando columnas innecesarias…")
        drop_cols = sorted(
            [idx for name, idx in col_name_to_idx.items() if name not in COLUMNS_TO_KEEP],
            reverse=True
        )

        for col_idx in drop_cols:
            ws.delete_cols(col_idx)

        # Eliminar filas superiores
        if header_row_1based > 1:
            progress("Eliminando filas de logo/título…")
            ws.delete_rows(1, header_row_1based - 1)

        # Reconstruir mapa de columnas (Encabezado en fila 1)
        col_name_to_new_idx = {}
        for cell in ws[1]:
            if cell.value:
                col_name_to_new_idx[str(cell.value).strip()] = cell.column

        remain_col_idx = col_name_to_new_idx.get('Remain time (H)')
        if remain_col_idx is None:
            raise ValueError("No se encontró 'Remain time (H)' tras la limpieza.")

        # Ajustar anchos
        progress("Ajustando anchos de columna…")
        ws.column_dimensions.clear()

        for col_name, col_idx in col_name_to_new_idx.items():
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = COL_WIDTHS.get(col_name, 15)

        # Altura y alineación
        progress("Ajustando formato y alineación…")
        ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

        total_rows = ws.max_row
        for row_idx in range(2, total_rows + 1):
            ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT

        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        data_alignment = Alignment(vertical="center", wrap_text=False)

        for cell in ws[1]:
            cell.alignment = header_alignment

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = data_alignment

        # Resaltar valores negativos
        progress("Resaltando valores negativos…")
        colored = 0

        for row in ws.iter_rows(min_row=2):
            cell = row[remain_col_idx - 1]
            try:
                if cell.value is not None and float(cell.value) < 0:
                    cell.fill = NEG_FILL
                    f = cell.font
                    cell.font = Font(
                        name=f.name,
                        size=f.size,
                        bold=True,
                        italic=f.italic,
                        underline=f.underline,
                        strike=f.strike,
                        color=NEG_COLOR
                    )
                    colored += 1
            except (ValueError, TypeError):
                pass

        # Autofiltro
        progress("Aplicando autofiltro…")
        last_col = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"

        # Guardar
        progress("Generando archivo temporal…")
        wb.save(out_path)
        wb.close()

        progress(f"¡Listo! — {colored} celdas negativas resaltadas")
        return colored