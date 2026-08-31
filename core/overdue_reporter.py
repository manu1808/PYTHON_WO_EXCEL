"""
Lógica para generación del reporte de Work Orders Vencidas, tablas resumen y gráficos en Excel.
"""

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

from config import HEADER_ROW_HEIGHT, DATA_ROW_HEIGHT, NEG_FILL, NEG_COLOR

class OverdueReporter:
    @staticmethod
    def create_overdue_report(clean_path: str, overdue_path: str, progress_cb=None) -> int:
        """
        Filtra las WOs vencidas (Remain time < 0) y crea un libro con resumen ejecutivo,
        tablas dinamizadas por CD/FT y gráficos de barras integrados.
        """
        def progress(msg):
            if progress_cb:
                progress_cb(msg)

        thin_border = Border(
            left=Side(style="thin", color="B7B7B7"),
            right=Side(style="thin", color="B7B7B7"),
            top=Side(style="thin", color="B7B7B7"),
            bottom=Side(style="thin", color="B7B7B7")
        )
        header_fill = PatternFill("solid", fgColor="D9EAD3")
        title_fill = PatternFill("solid", fgColor="1B4332")

        progress("Abriendo archivo limpio para crear reporte de vencidas…")
        wb = load_workbook(clean_path)
        ws = wb.active

        headers = {}
        for cell in ws[1]:
            if cell.value is not None:
                headers[str(cell.value).strip()] = cell.column

        required_columns = ["WO code", "CD", "FT", "Remain time (H)"]
        missing = [col for col in required_columns if col not in headers]
        if missing:
            wb.close()
            raise ValueError(f"No se encontraron las columnas necesarias: {missing}")

        cd_col = headers["CD"]
        ft_col = headers["FT"]
        remain_col = headers["Remain time (H)"]

        progress("Identificando filas vencidas…")
        overdue_rows = []
        cd_counts = {}
        ft_counts = {}

        for row_idx in range(2, ws.max_row + 1):
            remain_value = ws.cell(row=row_idx, column=remain_col).value
            is_overdue = False
            try:
                if remain_value is not None and float(remain_value) < 0:
                    is_overdue = True
            except (ValueError, TypeError):
                pass

            if not is_overdue:
                continue

            overdue_rows.append(row_idx)

            # Recuento CD
            cd = ws.cell(row=row_idx, column=cd_col).value
            cd = "(Sin CD)" if cd is None or str(cd).strip() == "" else str(cd).strip()
            cd_counts[cd] = cd_counts.get(cd, 0) + 1

            # Recuento FT
            ft = ws.cell(row=row_idx, column=ft_col).value
            ft = "(Sin FT)" if ft is None or str(ft).strip() == "" else str(ft).strip()
            ft_counts[ft] = ft_counts.get(ft, 0) + 1

        overdue_count = len(overdue_rows)

        # Ordenar recuentos
        cd_counts = dict(sorted(cd_counts.items(), key=lambda item: item[1], reverse=True))
        ft_counts = dict(sorted(ft_counts.items(), key=lambda item: item[1], reverse=True))

        # Eliminar no vencidas
        progress("Creando listado de vencidas…")
        overdue_set = set(overdue_rows)
        for row_idx in range(ws.max_row, 1, -1):
            if row_idx not in overdue_set:
                ws.delete_rows(row_idx)

        # Formato hoja Vencidas
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        for cell in ws[1]:
            cell.alignment = header_alignment
            orig_font = cell.font
            cell.font = Font(
                name=orig_font.name, size=orig_font.size, bold=True,
                italic=orig_font.italic, underline=orig_font.underline,
                strike=orig_font.strike, color=orig_font.color
            )

        ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

        current_headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}
        current_remain_col = current_headers["Remain time (H)"]

        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT
            cell = ws.cell(row=row_idx, column=current_remain_col)
            try:
                if cell.value is not None and float(cell.value) < 0:
                    cell.fill = NEG_FILL
                    orig_font = cell.font
                    cell.font = Font(
                        name=orig_font.name, size=orig_font.size, bold=True,
                        italic=orig_font.italic, underline=orig_font.underline,
                        strike=orig_font.strike, color=NEG_COLOR
                    )
            except (ValueError, TypeError):
                pass

        last_col = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
        ws.freeze_panes = "A2"

        # Crear hoja Resumen
        progress("Creando hoja de resumen…")
        ws_summary = wb.create_sheet("Resumen", 0)

        # Título General
        ws_summary.merge_cells("A1:H1")
        ws_summary["A1"] = "REPORTE DE WOs VENCIDAS"
        ws_summary["A1"].font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
        ws_summary["A1"].fill = title_fill
        ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary.row_dimensions[1].height = 30

        # Tabla CD
        ws_summary.merge_cells("A3:B3")
        ws_summary["A3"] = "RECUENTO DE WO CODE POR CD"
        ws_summary["A3"].font = Font(name="Segoe UI", size=12, bold=True)
        ws_summary["A3"].alignment = Alignment(horizontal="center", vertical="center")

        ws_summary["A4"] = "CD"
        ws_summary["B4"] = "Recuento de WO code"

        for cell in ws_summary[4]:
            if cell.column <= 2:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        cd_row = 5
        for cd, count in cd_counts.items():
            ws_summary.cell(row=cd_row, column=1, value=cd)
            ws_summary.cell(row=cd_row, column=2, value=count)
            for col in range(1, 3):
                cell = ws_summary.cell(row=cd_row, column=col)
                cell.border = thin_border
                if col == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            cd_row += 1

        cd_last_row = cd_row - 1

        # Gráfico CD
        if cd_counts:
            chart_cd = BarChart()
            chart_cd.type = "bar"
            chart_cd.style = 10
            chart_cd.title = "WOs vencidas por CD"
            chart_cd.y_axis.title = "CD"
            chart_cd.x_axis.title = "Cantidad de WOs"

            data_cd = Reference(ws_summary, min_col=2, min_row=4, max_row=cd_last_row)
            categories_cd = Reference(ws_summary, min_col=1, min_row=5, max_row=cd_last_row)
            chart_cd.add_data(data_cd, titles_from_data=True)
            chart_cd.set_categories(categories_cd)
            chart_cd.legend = None

            chart_cd.dataLabels = DataLabelList()
            chart_cd.dataLabels.showVal = True
            chart_cd.dataLabels.showLegendKey = False
            chart_cd.dataLabels.showCatName = False
            chart_cd.dataLabels.showSerName = False
            chart_cd.dataLabels.showPercent = False
            chart_cd.dataLabels.showLeaderLines = False

            chart_cd.height = 9
            chart_cd.width = 19
            chart_cd.gapWidth = 50
            ws_summary.add_chart(chart_cd, "D3")

        # Tabla FT
        ft_title_row = max(22, cd_last_row + 15)
        ws_summary.merge_cells(start_row=ft_title_row, start_column=1, end_row=ft_title_row, end_column=2)
        ws_summary.cell(row=ft_title_row, column=1, value="RECUENTO DE WO CODE POR FT")
        ws_summary.cell(row=ft_title_row, column=1).font = Font(name="Segoe UI", size=12, bold=True)
        ws_summary.cell(row=ft_title_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ft_header_row = ft_title_row + 1
        ws_summary.cell(row=ft_header_row, column=1, value="FT")
        ws_summary.cell(row=ft_header_row, column=2, value="Recuento de WO code")

        for col in range(1, 3):
            cell = ws_summary.cell(row=ft_header_row, column=col)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        ft_row = ft_header_row + 1
        for ft, count in ft_counts.items():
            ws_summary.cell(row=ft_row, column=1, value=ft)
            ws_summary.cell(row=ft_row, column=2, value=count)
            for col in range(1, 3):
                cell = ws_summary.cell(row=ft_row, column=col)
                cell.border = thin_border
                if col == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            ft_row += 1

        ft_last_row = ft_row - 1

        # Gráfico FT
        if ft_counts:
            chart_ft = BarChart()
            chart_ft.type = "bar"
            chart_ft.style = 10
            chart_ft.title = "WOs vencidas por FT"
            chart_ft.y_axis.title = "FT"
            chart_ft.x_axis.title = "Cantidad de WOs"

            data_ft = Reference(ws_summary, min_col=2, min_row=ft_header_row, max_row=ft_last_row)
            categories_ft = Reference(ws_summary, min_col=1, min_row=ft_header_row + 1, max_row=ft_last_row)
            chart_ft.add_data(data_ft, titles_from_data=True)
            chart_ft.set_categories(categories_ft)
            chart_ft.legend = None

            chart_ft.dataLabels = DataLabelList()
            chart_ft.dataLabels.showVal = True
            chart_ft.dataLabels.showLegendKey = False
            chart_ft.dataLabels.showCatName = False
            chart_ft.dataLabels.showSerName = False
            chart_ft.dataLabels.showPercent = False
            chart_ft.dataLabels.showLeaderLines = False

            chart_ft.height = 9
            chart_ft.width = 19
            chart_ft.gapWidth = 50
            ws_summary.add_chart(chart_ft, f"D{ft_title_row}")

        # Anchos de columna
        widths = {"A": 28, "B": 25, "C": 3, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18}
        for col, width in widths.items():
            ws_summary.column_dimensions[col].width = width

        progress("Guardando reporte de vencidas…")
        wb.save(overdue_path)
        wb.close()

        progress(f"Reporte de vencidas creado: {overdue_count} WOs")
        return overdue_count